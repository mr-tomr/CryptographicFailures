#!/usr/bin/env python3
"""
CryptographicFailures_Automated.py

Consolidates the CryptographicFailures repository workflow into one tool:
  * single host + single port
  * single host + ports file
  * hosts file + single port
  * hosts file + ports file
  * hosts file containing host:port pairs
  * existing standard-format Nmap (.nmap) scan import; detects SSL/TLS host:port targets

Outputs:
  * raw Nmap evidence per target
  * classified per-target cipher files
  * report-ready tabbed TXT (host:port + indented cipher findings)
  * CSV with host/IP and port in separate columns; one cipher finding per row
  * DOCX with each host:port as a real Word Heading 4 paragraph
  * SWEET32 TXT / CSV / DOCX outputs

Manual cleanup previously documented in DisplayGroupOfFindings.sh is automated:
  * deduplication
  * blank-result suppression
  * exact weak/insecure filtering (does NOT accidentally drop 'insecure')
  * SWEET32 extraction
  * instance/finding counts
  * filename cleanup and host:port formatting

Requires:
  * nmap
  * Python 3
  * python-docx (pip install python-docx)

Cipher classifications are downloaded from ciphersuite.info once per run.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import ipaddress
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    Document = None

API_BASE = "https://ciphersuite.info/api/cs/security"
CLASSIFICATIONS = ("insecure", "weak", "secure", "recommended")
REPORTABLE = {"insecure", "weak"}
SWEET32_RE = re.compile(r"(?:3DES_EDE_CBC|DES_CBC_SHA)", re.IGNORECASE)
CIPHER_RE = re.compile(r"\b(TLS_[A-Za-z0-9_]+)\b")
PORT_RE = re.compile(r"^[0-9]{1,5}$")
NMAP_REPORT_RE = re.compile(r"^Nmap scan report for\s+(.+?)\s*$")
NMAP_OPEN_PORT_RE = re.compile(r"^(\d+)/tcp\s+open\b")


@dataclass(frozen=True, order=True)
class Target:
    host: str
    port: int

    @property
    def display(self) -> str:
        # Word/report format required by the workflow.
        if ":" in self.host and not self.host.startswith("["):
            return f"[{self.host}]:{self.port}"
        return f"{self.host}:{self.port}"

    @property
    def file_stem(self) -> str:
        # Legacy-style underscore separator for filesystem artifacts only.
        safe_host = re.sub(r"[^A-Za-z0-9._-]+", "-", self.host.strip("[]"))
        return f"{safe_host}_{self.port}"


@dataclass(frozen=True)
class Finding:
    target: Target
    cipher: str
    classification: str

    @property
    def text(self) -> str:
        # User requested classification kept in same cell/string as the cipher.
        return f"{self.cipher}: {self.classification}"


def die(message: str, code: int = 1) -> None:
    print(f"[!] {message}", file=sys.stderr)
    raise SystemExit(code)


def validate_port(value: str | int) -> int:
    s = str(value).strip()
    if not PORT_RE.fullmatch(s):
        raise ValueError(f"Invalid port: {value}")
    port = int(s)
    if not 1 <= port <= 65535:
        raise ValueError(f"Port out of range: {port}")
    return port


def clean_lines(path: Path) -> List[str]:
    if not path.is_file():
        raise FileNotFoundError(path)
    result: List[str] = []
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        result.append(line)
    return result


def normalize_host(host: str) -> str:
    host = host.strip()
    if not host:
        raise ValueError("Host cannot be blank")
    if host.startswith("[") and host.endswith("]"):
        host = host[1:-1]
    if any(ch.isspace() for ch in host):
        raise ValueError(f"Invalid host containing whitespace: {host!r}")
    return host


def parse_host_port(value: str) -> Target:
    value = value.strip()
    if not value:
        raise ValueError("Blank host:port entry")

    # [IPv6]:443
    m = re.fullmatch(r"\[(.+)]:(\d{1,5})", value)
    if m:
        return Target(normalize_host(m.group(1)), validate_port(m.group(2)))

    # hostname:443 or IPv4:443. For raw IPv6, require [IPv6]:port to avoid ambiguity.
    if value.count(":") == 1:
        host, port = value.rsplit(":", 1)
        return Target(normalize_host(host), validate_port(port))

    if value.count(":") > 1:
        raise ValueError(
            f"Ambiguous host:port entry {value!r}. Use [IPv6-address]:port for IPv6."
        )
    raise ValueError(f"Expected host:port, got: {value!r}")



def parse_nmap_report_host(value: str) -> str:
    """Return the best scan target from an Nmap scan-report line."""
    value = value.strip()
    m = re.fullmatch(r".+\s+\(([^()]+)\)", value)
    if m:
        return normalize_host(m.group(1))
    return normalize_host(value)


def parse_nmap_ssl_targets(nmap_path: Path) -> List[Target]:
    """
    Parse standard Nmap text output and return host:port pairs where
    ssl-enum-ciphers produced output. Multi-host files retain host association.
    """
    if not nmap_path.is_file():
        raise FileNotFoundError(nmap_path)

    targets: List[Target] = []
    current_host: str | None = None
    current_port: int | None = None
    port_has_ssl_enum = False

    def finish_port() -> None:
        nonlocal current_port, port_has_ssl_enum
        if current_host is not None and current_port is not None and port_has_ssl_enum:
            targets.append(Target(current_host, current_port))
        current_port = None
        port_has_ssl_enum = False

    for raw in nmap_path.read_text(encoding="utf-8", errors="replace").splitlines():
        line = raw.strip()

        report = NMAP_REPORT_RE.match(line)
        if report:
            finish_port()
            current_host = parse_nmap_report_host(report.group(1))
            continue

        port_match = NMAP_OPEN_PORT_RE.match(line)
        if port_match:
            finish_port()
            current_port = validate_port(port_match.group(1))
            continue

        if current_port is not None and "ssl-enum-ciphers:" in line:
            port_has_ssl_enum = True

    finish_port()
    return dedupe_targets(targets)


def dedupe_targets(targets: Iterable[Target]) -> List[Target]:
    return sorted(set(targets), key=lambda t: (host_sort_key(t.host), t.port))


def host_sort_key(host: str):
    try:
        ip = ipaddress.ip_address(host)
        return (0, ip.version, int(ip))
    except ValueError:
        return (1, host.lower())


def build_targets(mode: int, host: str | None, port: int | None,
                  hosts_file: Path | None, ports_file: Path | None,
                  pairs_file: Path | None) -> List[Target]:
    targets: List[Target] = []

    if mode == 1:
        if host is None or port is None:
            raise ValueError("Mode 1 requires host and port")
        targets = [Target(normalize_host(host), validate_port(port))]

    elif mode == 2:
        if host is None or ports_file is None:
            raise ValueError("Mode 2 requires host and ports file")
        h = normalize_host(host)
        for p in clean_lines(ports_file):
            targets.append(Target(h, validate_port(p)))

    elif mode == 3:
        if hosts_file is None or port is None:
            raise ValueError("Mode 3 requires hosts file and port")
        p = validate_port(port)
        for h in clean_lines(hosts_file):
            targets.append(Target(normalize_host(h), p))

    elif mode == 4:
        if hosts_file is None or ports_file is None:
            raise ValueError("Mode 4 requires hosts file and ports file")
        hosts = [normalize_host(h) for h in clean_lines(hosts_file)]
        ports = [validate_port(p) for p in clean_lines(ports_file)]
        for h in hosts:
            for p in ports:
                targets.append(Target(h, p))

    elif mode == 5:
        if pairs_file is None:
            raise ValueError("Mode 5 requires a host:port pairs file")
        for line in clean_lines(pairs_file):
            targets.append(parse_host_port(line))

    else:
        raise ValueError(f"Unknown mode: {mode}")

    targets = dedupe_targets(targets)
    if not targets:
        raise ValueError("No scan targets were produced from the supplied input")
    return targets


def prompt_existing_file(label: str) -> Path:
    while True:
        value = input(f"{label}: ").strip()
        p = Path(os.path.expanduser(value))
        if p.is_file():
            return p
        print(f"[!] File not found: {p}")


def prompt_port(label: str = "Port") -> int:
    while True:
        try:
            return validate_port(input(f"{label}: ").strip())
        except ValueError as exc:
            print(f"[!] {exc}")


def interactive_targets() -> List[Target]:
    print("\nCryptographic Failures - Scan and Report")
    print("----------------------------------------")
    print("1) Single host + single port")
    print("2) Single host + ports file")
    print("3) Hosts file + single port")
    print("4) Hosts file + ports file")
    print("5) Host:port pairs file")
    print("6) Import existing Nmap scan and detect SSL/TLS targets")
    print("0) Exit")

    while True:
        choice = input("\nSelection: ").strip()
        if choice == "0":
            raise SystemExit(0)
        if choice in {"1", "2", "3", "4", "5", "6"}:
            mode = int(choice)
            break
        print("[!] Choose 0-6.")

    if mode == 1:
        host = normalize_host(input("Host/IP: "))
        return build_targets(1, host, prompt_port(), None, None, None)
    if mode == 2:
        host = normalize_host(input("Host/IP: "))
        return build_targets(2, host, None, None, prompt_existing_file("Ports file"), None)
    if mode == 3:
        return build_targets(3, None, prompt_port(), prompt_existing_file("Hosts file"), None, None)
    if mode == 4:
        return build_targets(
            4, None, None,
            prompt_existing_file("Hosts file"),
            prompt_existing_file("Ports file"),
            None,
        )
    if mode == 5:
        return build_targets(5, None, None, None, None, prompt_existing_file("Host:port pairs file"))

    nmap_file = prompt_existing_file("Existing Nmap .nmap file")
    targets = parse_nmap_ssl_targets(nmap_file)
    if not targets:
        raise ValueError(
            "No SSL/TLS targets were detected. The imported Nmap file must contain "
            "ssl-enum-ciphers script output."
        )
    print(f"[+] Detected {len(targets)} SSL/TLS host:port target(s)")
    return targets


def fetch_classification_texts(cache_dir: Path) -> Dict[str, str]:
    cache_dir.mkdir(parents=True, exist_ok=True)
    result: Dict[str, str] = {}
    headers = {"Accept": "application/json", "User-Agent": "CryptographicFailures-Automated/1.0"}

    for classification in CLASSIFICATIONS:
        url = f"{API_BASE}/{classification}"
        print(f"[*] Downloading cipher classification: {classification}")
        req = urllib.request.Request(url, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=30) as response:
                text = response.read().decode("utf-8", errors="replace")
        except (urllib.error.URLError, TimeoutError) as exc:
            cached = cache_dir / f"{classification}.json"
            if cached.is_file():
                print(f"[!] API unavailable for {classification}; using cached {cached}")
                text = cached.read_text(encoding="utf-8", errors="replace")
            else:
                raise RuntimeError(f"Unable to download {url}: {exc}") from exc

        # Basic JSON validation, while retaining raw text for exact-name membership checks.
        try:
            json.loads(text)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"Invalid JSON received for {classification}: {exc}") from exc

        (cache_dir / f"{classification}.json").write_text(text, encoding="utf-8")
        result[classification] = text
    return result


def classify_cipher(cipher: str, classification_texts: Dict[str, str]) -> str:
    needle = f'"{cipher}"'
    matches = [name for name in CLASSIFICATIONS if needle in classification_texts.get(name, "")]
    if not matches:
        return "unclassified"
    # Prefer the more concerning category if an upstream dataset ever overlaps.
    for priority in ("insecure", "weak", "secure", "recommended"):
        if priority in matches:
            return priority
    return matches[0]


def run_nmap(target: Target, raw_dir: Path, timeout: int = 300) -> Path:
    raw_dir.mkdir(parents=True, exist_ok=True)
    outfile = raw_dir / f"{target.file_stem}.nmap"
    cmd = [
        "nmap", "-Pn",
        "--script", "ssl-cert,ssl-enum-ciphers",
        "-p", str(target.port),
        target.host,
        "-oN", str(outfile),
    ]
    print(f"[*] Scanning {target.display}")
    try:
        proc = subprocess.run(cmd, text=True, capture_output=True, timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"Nmap timed out for {target.display}") from exc
    if proc.returncode != 0:
        raise RuntimeError(
            f"Nmap failed for {target.display} (exit {proc.returncode}):\n{proc.stderr.strip()}"
        )
    return outfile


def parse_nmap_ciphers(nmap_path: Path) -> List[str]:
    text = nmap_path.read_text(encoding="utf-8", errors="replace")
    # Only cipher-suite names are extracted; duplicate protocol blocks collapse naturally.
    return sorted(set(CIPHER_RE.findall(text)))


def write_per_target(target: Target, findings: Sequence[Finding], target_dir: Path) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / f"{target.file_stem}.txt"
    rank = {"insecure": 0, "weak": 1, "secure": 2, "recommended": 3, "unclassified": 4}
    ordered = sorted(findings, key=lambda f: (rank.get(f.classification, 99), f.cipher))
    with path.open("w", encoding="utf-8") as fh:
        for finding in ordered:
            fh.write(f"{finding.text}\n")
    return path


def reportable_findings(findings: Sequence[Finding]) -> List[Finding]:
    return [f for f in findings if f.classification in REPORTABLE]


def sweet32_findings(findings: Sequence[Finding]) -> List[Finding]:
    return [f for f in findings if SWEET32_RE.search(f.cipher)]


def group_findings(findings: Sequence[Finding]) -> List[Tuple[Target, List[Finding]]]:
    grouped: Dict[Target, List[Finding]] = {}
    for finding in findings:
        grouped.setdefault(finding.target, []).append(finding)
    result = []
    for target in sorted(grouped, key=lambda t: (host_sort_key(t.host), t.port)):
        unique = {(f.cipher, f.classification): f for f in grouped[target]}
        rows = sorted(unique.values(), key=lambda f: (f.classification, f.cipher))
        result.append((target, rows))
    return result


def write_tabbed_txt(path: Path, findings: Sequence[Finding]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as fh:
        for target, rows in group_findings(findings):
            # No 'File:' prefix, no '.txt', and colon separator in final output.
            fh.write(f"{target.display}\n")
            for finding in rows:
                fh.write(f"\t{finding.text}\n")
            fh.write("\n")


def write_csv(path: Path, findings: Sequence[Finding]) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["IP/Host", "Port", "Cipher Finding"])
        for target, rows in group_findings(findings):
            for finding in rows:
                # Classification intentionally remains in the same cell as cipher.
                writer.writerow([target.host, target.port, finding.text])


def write_docx(path: Path, title: str, findings: Sequence[Finding], heading_level: int = 4) -> None:
    if Document is None:
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx")

    doc = Document()
    doc.add_heading(title, level=1)
    for target, rows in group_findings(findings):
        doc.add_heading(target.display, level=heading_level)
        for finding in rows:
            p = doc.add_paragraph(finding.text)
            p.paragraph_format.left_indent = Inches(0.25)
    doc.save(path)


def write_summary(path: Path, targets: Sequence[Target], all_findings: Sequence[Finding],
                  reportable: Sequence[Finding], sweet32: Sequence[Finding], errors: Sequence[str]) -> None:
    report_groups = group_findings(reportable)
    sweet_groups = group_findings(sweet32)
    scanned_targets = {f.target for f in all_findings}
    with path.open("w", encoding="utf-8") as fh:
        fh.write("Cryptographic Failures Scan Summary\n")
        fh.write("==================================\n")
        fh.write(f"Requested targets: {len(targets)}\n")
        fh.write(f"Targets with classified cipher data: {len(scanned_targets)}\n")
        fh.write(f"Reportable instances (host:port): {len(report_groups)}\n")
        fh.write(f"Reportable cipher rows: {len(reportable)}\n")
        fh.write(f"SWEET32 instances (host:port): {len(sweet_groups)}\n")
        fh.write(f"SWEET32 cipher rows: {len(sweet32)}\n")
        fh.write(f"Scan errors: {len(errors)}\n")
        if errors:
            fh.write("\nErrors\n------\n")
            for error in errors:
                fh.write(f"- {error}\n")


def create_outputs(output_dir: Path, targets: Sequence[Target], all_findings: Sequence[Finding],
                   errors: Sequence[str]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    reportable = reportable_findings(all_findings)
    sweet32 = sweet32_findings(all_findings)

    write_tabbed_txt(output_dir / "findings_tabbed.txt", reportable)
    write_csv(output_dir / "findings.csv", reportable)
    write_docx(output_dir / "findings.docx", "Cryptographic Failures - Instances", reportable, 4)

    write_tabbed_txt(output_dir / "sweet32_tabbed.txt", sweet32)
    write_csv(output_dir / "sweet32.csv", sweet32)
    write_docx(output_dir / "sweet32.docx", "SWEET32 - Instances", sweet32, 4)

    write_summary(output_dir / "summary.txt", targets, all_findings, reportable, sweet32, errors)


def scan_targets(targets: Sequence[Target], output_dir: Path, keep_going: bool = True) -> Tuple[List[Finding], List[str]]:
    if shutil.which("nmap") is None:
        die("nmap was not found in PATH. Install nmap before scanning.")

    raw_dir = output_dir / "raw_nmap"
    target_dir = output_dir / "classified_targets"
    cache_dir = output_dir / "cipher_api_cache"
    classification_texts = fetch_classification_texts(cache_dir)

    all_findings: List[Finding] = []
    errors: List[str] = []

    for index, target in enumerate(targets, 1):
        print(f"\n[{index}/{len(targets)}] {target.display}")
        try:
            nmap_file = run_nmap(target, raw_dir)
            ciphers = parse_nmap_ciphers(nmap_file)
            findings = [
                Finding(target, cipher, classify_cipher(cipher, classification_texts))
                for cipher in ciphers
            ]
            # Deduplication is performed before writing.
            findings = list({(f.cipher, f.classification): f for f in findings}.values())
            write_per_target(target, findings, target_dir)
            all_findings.extend(findings)
            report_count = len(reportable_findings(findings))
            print(f"[+] {len(ciphers)} unique cipher(s), {report_count} reportable")
        except Exception as exc:  # keep batch scans useful while preserving error evidence in summary
            message = f"{target.display}: {exc}"
            errors.append(message)
            print(f"[!] {message}")
            if not keep_going:
                raise

    return all_findings, errors


def default_output_dir() -> Path:
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path.cwd() / f"cryptographic_failures_{stamp}"


def self_test() -> int:
    """Offline test of parsing, filtering, CSV/TXT/DOCX generation."""
    test_dir = Path.cwd() / "cryptographic_failures_selftest"
    if test_dir.exists():
        shutil.rmtree(test_dir)
    test_dir.mkdir(parents=True)

    target1 = Target("10.10.10.10", 443)
    target2 = Target("server.example.com", 8443)

    imported_nmap = test_dir / "import_test.nmap"
    imported_nmap.write_text(
        """Nmap scan report for web1.example.com (10.10.10.10)
443/tcp open  https
| ssl-enum-ciphers:
|_  least strength: C
80/tcp open http
Nmap scan report for server.example.com
8443/tcp open  https-alt
| ssl-enum-ciphers:
|_  least strength: B
22/tcp open ssh
""",
        encoding="utf-8",
    )
    imported_targets = parse_nmap_ssl_targets(imported_nmap)
    assert imported_targets == [target1, target2]
    findings = [
        Finding(target1, "TLS_RSA_WITH_3DES_EDE_CBC_SHA", "insecure"),
        Finding(target1, "TLS_RSA_WITH_AES_128_CBC_SHA", "weak"),
        Finding(target1, "TLS_AES_256_GCM_SHA384", "recommended"),
        # duplicate validates dedupe in grouped outputs
        Finding(target1, "TLS_RSA_WITH_AES_128_CBC_SHA", "weak"),
        Finding(target2, "TLS_ECDHE_RSA_WITH_AES_128_CBC_SHA", "weak"),
        Finding(target2, "TLS_AES_128_GCM_SHA256", "secure"),
    ]
    create_outputs(test_dir, [target1, target2], findings, [])

    tabbed = (test_dir / "findings_tabbed.txt").read_text(encoding="utf-8")
    assert "File:" not in tabbed
    assert ".txt" not in tabbed
    assert "10.10.10.10:443" in tabbed
    assert "\tTLS_RSA_WITH_3DES_EDE_CBC_SHA: insecure" in tabbed
    assert "recommended" not in tabbed
    assert tabbed.count("TLS_RSA_WITH_AES_128_CBC_SHA: weak") == 1

    with (test_dir / "findings.csv").open(newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == ["IP/Host", "Port", "Cipher Finding"]
    assert ["10.10.10.10", "443", "TLS_RSA_WITH_3DES_EDE_CBC_SHA: insecure"] in rows

    print(f"[+] Self-test passed: {test_dir}")
    return 0


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Scan TLS cipher suites and automatically generate report-ready TXT, CSV, and DOCX outputs."
    )
    p.add_argument("--mode", type=int, choices=range(1, 7), help="Input mode 1-6; omit for interactive menu")
    p.add_argument("--host", help="Host/IP for modes 1-2")
    p.add_argument("--port", type=int, help="Port for modes 1 or 3")
    p.add_argument("--hosts-file", type=Path, help="Hosts file for modes 3-4")
    p.add_argument("--ports-file", type=Path, help="Ports file for modes 2 or 4")
    p.add_argument("--pairs-file", type=Path, help="host:port pairs file for mode 5")
    p.add_argument("--nmap-file", type=Path, help="existing standard-format Nmap .nmap file for mode 6")
    p.add_argument("--output-dir", type=Path, help="Output directory; default is timestamped")
    p.add_argument("--stop-on-error", action="store_true", help="Stop batch on first target scan error")
    p.add_argument("--self-test", action="store_true", help="Run offline parser/report self-test and exit")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    if args.self_test:
        return self_test()

    try:
        if args.mode is None:
            targets = interactive_targets()
        elif args.mode == 6:
            if args.nmap_file is None:
                raise ValueError("Mode 6 requires --nmap-file")
            targets = parse_nmap_ssl_targets(args.nmap_file.expanduser())
            if not targets:
                raise ValueError(
                    "No SSL/TLS targets were detected. The imported Nmap file must contain "
                    "ssl-enum-ciphers script output."
                )
        else:
            targets = build_targets(
                args.mode, args.host, args.port,
                args.hosts_file, args.ports_file, args.pairs_file,
            )
    except (ValueError, FileNotFoundError) as exc:
        die(str(exc))

    output_dir = (args.output_dir or default_output_dir()).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"\n[*] Targets: {len(targets)}")
    print(f"[*] Output:  {output_dir}")
    for target in targets:
        print(f"    {target.display}")

    try:
        all_findings, errors = scan_targets(targets, output_dir, keep_going=not args.stop_on_error)
        create_outputs(output_dir, targets, all_findings, errors)
    except Exception as exc:
        die(str(exc))

    reportable = reportable_findings(all_findings)
    print("\n[+] Complete")
    print(f"[+] Reportable host:port instances: {len(group_findings(reportable))}")
    print(f"[+] Reportable cipher rows:          {len(reportable)}")
    print(f"[+] Results directory:               {output_dir}")
    print("[+] Word output:                     findings.docx (host:port = Heading 4)")
    if errors:
        print(f"[!] {len(errors)} target(s) had scan errors; see summary.txt")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
